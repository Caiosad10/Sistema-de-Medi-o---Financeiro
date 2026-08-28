import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy
from paths import get_data_path, get_resource_path
import os

# ===== CORREÇÃO DOS CAMINHOS =====
DATA_PATH = get_data_path()
DB_NAME = os.path.join(DATA_PATH, "medicao.db")
EXCEL_DADOS = "PROJETO_AEV_CORRIGIDO_COM_DADOS.xlsx"
EXCEL_TEMPLATE = "TEMPLATE_MEDICAO_AEV_LIMPO_V2_VALIDADO.xlsx"

PROJETOS_SISTEMA = [
    "PROJETO AEV", "REUSO", "REUSO MW", "REMANEJAMENTO", 
    "TIM DISTRATO", "TIM ROLLOUT", "TIM MOU", "TIM DROP", 
    "FIBER", "BOOTS", "BASE UNICA", "DESMOBILIZAÇÃO", 
    "LOG REVERSA", "ANJLO", "PROJETO BPR", "REDE EXTERNA", "MEDIDORES"
]

COLUNAS_HISTORICO_EXTRAS = [
    ("atividade", "TEXT"),
    ("valor_total", "REAL DEFAULT 0.0"),
    ("valor_liquido", "REAL DEFAULT 0.0"),
    ("oc", "TEXT"),
    ("end_id", "TEXT"),
    ("nota_fiscal", "TEXT"),
    ("regional", "TEXT"),
    ("tipo_ciclo", "TEXT"),
]

# Tipos de ciclo (abas de medição)
TIPO_NOVA_MEDICAO = "NOVA_MEDICAO"
TIPO_SALDOS_FINAIS = "SALDOS_FINAIS"
TIPO_SOLICITACAO_EMERGENCIAL = "SOLICITACAO_EMERGENCIAL"

TIPOS_CICLO = {
    TIPO_NOVA_MEDICAO: "Medição",
    TIPO_SALDOS_FINAIS: "Saldos Finais",
    TIPO_SOLICITACAO_EMERGENCIAL: "Solicitação Emergencial",
}

LIMITE_ABAS_POR_PROJETO = 8

SELECT_LANCAMENTO_COLUNAS = """
    l.pc AS "PC", l.prestador AS "PRESTADOR", l.atividade AS "ATIVIDADE/SERVIÇO",
    l.valor_total AS "VALOR TOTAL DO PC", l.valor_liquido AS "VALOR LIQUIDO",
    l.valores_pagos AS "VALORES PAGOS", l.saldo_a_pagar AS "SALDO A PAGAR",
    l.valor_a_pagar AS "VALOR A PAGAR", l.saldo_aberto AS "SALDO EM ABERTO",
    l.status_pc AS "STATUS PC", l.oc AS "OC", l.end_id AS "END ID",
    l.nota_fiscal AS "NOTA FISCAL", l.regional AS "REGIONAL",
    l.ciclo_id AS "CICLO_ID"
"""

def get_connection():
    # Cria a pasta se não existir
    os.makedirs(DATA_PATH, exist_ok=True)
    conn = sqlite3.connect(DB_NAME, timeout=10.0)
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def _migrar_banco(cursor):
    """Adiciona colunas/tabelas novas em bancos já existentes sem perder dados."""
    # --- histórico: colunas extras ---
    cursor.execute("PRAGMA table_info(historico_medicoes)")
    colunas_hist = {row[1] for row in cursor.fetchall()}
    for nome, tipo in COLUNAS_HISTORICO_EXTRAS:
        if nome not in colunas_hist:
            cursor.execute(f"ALTER TABLE historico_medicoes ADD COLUMN {nome} {tipo}")

    # --- tabela ciclos (abas de medição) ---
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS ciclos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        projeto_id INTEGER NOT NULL,
        titulo TEXT,
        tipo TEXT NOT NULL DEFAULT 'NOVA_MEDICAO',
        aberto INTEGER NOT NULL DEFAULT 1,
        data_criacao TEXT,
        ordem INTEGER DEFAULT 0,
        FOREIGN KEY (projeto_id) REFERENCES projetos(id) ON DELETE CASCADE
    )
    """)

    # --- lancamentos: coluna ciclo_id ---
    cursor.execute("PRAGMA table_info(lancamentos)")
    colunas_lanc = {row[1] for row in cursor.fetchall()}
    if "ciclo_id" not in colunas_lanc:
        cursor.execute("ALTER TABLE lancamentos ADD COLUMN ciclo_id INTEGER REFERENCES ciclos(id)")

    # --- migração: lançamentos sem ciclo → um ciclo "Medição" por projeto ---
    # Se o projeto só tem residual (valor_a_pagar = 0), o ciclo nasce FECHADO,
    # para não reaparecer como "medição ativa" na tela de Lançamentos.
    cursor.execute("""
        SELECT DISTINCT l.projeto_id
        FROM lancamentos l
        WHERE l.ciclo_id IS NULL
    """)
    projetos_sem_ciclo = [row[0] for row in cursor.fetchall()]
    from datetime import datetime
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")
    for projeto_id in projetos_sem_ciclo:
        cursor.execute("""
            SELECT COUNT(*) FROM lancamentos
            WHERE projeto_id = ? AND ciclo_id IS NULL AND valor_a_pagar > 0
        """, (projeto_id,))
        tem_ativo = cursor.fetchone()[0] > 0
        aberto = 1 if tem_ativo else 0
        cursor.execute("""
            INSERT INTO ciclos (projeto_id, titulo, tipo, aberto, data_criacao, ordem)
            VALUES (?, ?, ?, ?, ?, 0)
        """, (projeto_id, "Medição", TIPO_NOVA_MEDICAO, aberto, agora))
        ciclo_id = cursor.lastrowid
        cursor.execute("""
            UPDATE lancamentos SET ciclo_id = ?
            WHERE projeto_id = ? AND ciclo_id IS NULL
        """, (ciclo_id, projeto_id))

    # Ciclos abertos que só têm residual (valor_a_pagar = 0) e já têm linhas
    # → fecha, para não poluir a tela após migração da versão antiga.
    # Abas vazias (sem lançamentos) continuam abertas.
    cursor.execute("""
        SELECT c.id
        FROM ciclos c
        WHERE c.aberto = 1
          AND EXISTS (
              SELECT 1 FROM lancamentos l WHERE l.ciclo_id = c.id
          )
          AND NOT EXISTS (
              SELECT 1 FROM lancamentos l
              WHERE l.ciclo_id = c.id AND l.valor_a_pagar > 0
          )
    """)
    for (cid,) in cursor.fetchall():
        cursor.execute("UPDATE ciclos SET aberto = 0 WHERE id = ?", (cid,))

def inicializar_banco():
    """Inicializa o banco de dados com as tabelas e dados iniciais"""
    with get_connection() as conn:
        cursor = conn.cursor()
        
        # Criação das tabelas projetos e fornecedores mantida...
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS projetos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT UNIQUE NOT NULL
        )
        """)
        
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS fornecedores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            razao_social TEXT UNIQUE NOT NULL,
            razao_fantasia TEXT,
            cnpj_cpf TEXT UNIQUE NOT NULL,
            banco TEXT,
            agencia TEXT,
            conta TEXT,
            tipo_pix TEXT,
            chave_pix TEXT
        )
        """)

        cursor.execute("""
        CREATE TABLE IF NOT EXISTS ciclos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            projeto_id INTEGER NOT NULL,
            titulo TEXT,
            tipo TEXT NOT NULL DEFAULT 'NOVA_MEDICAO',
            aberto INTEGER NOT NULL DEFAULT 1,
            data_criacao TEXT,
            ordem INTEGER DEFAULT 0,
            FOREIGN KEY (projeto_id) REFERENCES projetos(id) ON DELETE CASCADE
        )
        """)
        
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS lancamentos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            projeto_id INTEGER NOT NULL,
            ciclo_id INTEGER,
            pc TEXT NOT NULL,
            prestador TEXT,
            atividade TEXT,
            valor_total REAL DEFAULT 0.0 CHECK(valor_total >= 0),
            valor_liquido REAL DEFAULT 0.0,
            valores_pagos REAL DEFAULT 0.0 CHECK(valores_pagos >= 0),
            saldo_a_pagar REAL DEFAULT 0.0,
            valor_a_pagar REAL DEFAULT 0.0 CHECK(valor_a_pagar >= 0),
            saldo_aberto REAL DEFAULT 0.0,
            status_pc TEXT,
            oc TEXT,
            end_id TEXT,
            nota_fiscal TEXT,
            regional TEXT,
            FOREIGN KEY (projeto_id) REFERENCES projetos(id) ON DELETE CASCADE,
            FOREIGN KEY (ciclo_id) REFERENCES ciclos(id) ON DELETE SET NULL
        )
        """)

        cursor.execute("""
        CREATE TABLE IF NOT EXISTS historico_medicoes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            projeto_id INTEGER NOT NULL,
            data_fechamento TEXT NOT NULL,
            mes_referencia TEXT NOT NULL,
            pc TEXT NOT NULL,
            prestador TEXT,
            atividade TEXT,
            valor_total REAL DEFAULT 0.0,
            valor_liquido REAL DEFAULT 0.0,
            valor_medido REAL,
            valores_pagos REAL,
            saldo_a_pagar REAL,
            valor_a_pagar REAL,
            saldo_aberto REAL,
            status_pc TEXT,
            oc TEXT,
            end_id TEXT,
            nota_fiscal TEXT,
            regional TEXT,
            status TEXT,
            tipo_ciclo TEXT,
            data_registro DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (projeto_id) REFERENCES projetos(id) ON DELETE CASCADE
        )
        """)
        
        _migrar_banco(cursor)

        for proj in PROJETOS_SISTEMA:
            cursor.execute("INSERT OR IGNORE INTO projetos (nome) VALUES (?)", (proj,))

        # Importa planilha base SOMENTE em instalação virgem.
        # Se já existe histórico, fornecedores ou lançamentos, o banco já foi usado
        # e NÃO deve ser repovoado (mesmo com zero lançamentos abertos).
        cursor.execute("SELECT COUNT(*) FROM lancamentos")
        qtd_lanc = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM historico_medicoes")
        qtd_hist = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM fornecedores")
        qtd_forn = cursor.fetchone()[0]

        if qtd_lanc == 0 and qtd_hist == 0 and qtd_forn == 0:
            importar_dados_iniciais(conn)
            # Garante que lançamentos importados recebam um ciclo
            _migrar_banco(cursor)
        
        conn.commit()

def _obter_aba(wb, nomes_possiveis):
    for nome in nomes_possiveis:
        if nome in wb.sheetnames:
            return wb[nome]
    return None

def _tratar_dados_lancamento(dados):
    """Higieniza e valida os tipos de dados do lançamento antes de interagir com o SQLite."""
    lista = list(dados[:14])
    while len(lista) < 14:
        lista.append("")
        
    # Colunas numéricas (valor_total, valor_liquido, valores_pagos, saldo_a_pagar, valor_a_pagar, saldo_aberto)
    for i in range(3, 9):
        try:
            val = float(lista[i]) if lista[i] not in (None, "") else 0.0
            lista[i] = max(0.0, val) if i in (3, 5, 7) else val # Impede valores negativos em total, pagos e a pagar
        except (ValueError, TypeError):
            lista[i] = 0.0
            
    # Colunas textuais
    for i in [0, 1, 2, 9, 10, 11, 12, 13]:
        lista[i] = str(lista[i]).strip() if lista[i] is not None else ""
        
    return lista

def importar_dados_iniciais(conn):
    """Importa dados iniciais dos arquivos Excel - CORRIGIDO com verificação de existência"""
    try:
        # Tenta carregar o arquivo de dados
        dados_path = get_resource_path(EXCEL_DADOS)
        
        if not os.path.exists(dados_path):
            print(f"Arquivo {EXCEL_DADOS} não encontrado. Importação inicial ignorada.")
            return
        
        wb = openpyxl.load_workbook(dados_path, data_only=True)
        cursor = conn.cursor()
        
        cursor.execute("SELECT id FROM projetos WHERE nome = 'PROJETO AEV'")
        res = cursor.fetchone()
        projeto_id = res[0] if res else 1

        # Importa fornecedores
        ws_cad = _obter_aba(wb, ["CADASTRO", "CADASTROS DE FORNECEDORES"])
        if ws_cad:
            for row in ws_cad.iter_rows(min_row=2, values_only=True):
                if row[0] and row[2]: 
                    cursor.execute("""
                    INSERT OR IGNORE INTO fornecedores 
                    (razao_social, razao_fantasia, cnpj_cpf, banco, agencia, conta, tipo_pix, chave_pix)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        str(row[0]).strip(),
                        str(row[1]).strip() if row[1] else "",
                        str(row[2]).strip(),
                        str(row[3]).strip() if row[3] else "",
                        str(row[4]).strip() if row[4] else "",
                        str(row[5]).strip() if row[5] else "",
                        str(row[6]).strip() if row[6] else "",
                        str(row[7]).strip() if row[7] else ""
                    ))

        # Importa lançamentos
        ws_lanc = _obter_aba(wb, ["LANÇAMENTOS", "27 07 2026"])
        if ws_lanc:
            min_r = 3 if "27 07 2026" in wb.sheetnames else 2
            for row in ws_lanc.iter_rows(min_row=min_r, values_only=True):
                if row[0] and not str(row[0]).startswith("PROJETO"): 
                    cursor.execute("""
                    INSERT INTO lancamentos (
                        projeto_id, pc, prestador, atividade, valor_total, valor_liquido,
                        valores_pagos, saldo_a_pagar, valor_a_pagar, saldo_aberto,
                        status_pc, oc, end_id, nota_fiscal, regional
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        projeto_id,
                        str(row[0]).strip(),
                        str(row[1]).strip() if row[1] else "",
                        str(row[2]).strip() if row[2] else "",
                        float(row[3] or 0.0),
                        float(row[4] or 0.0),
                        float(row[5] or 0.0),
                        float(row[6] or 0.0),
                        float(row[7] or 0.0),
                        float(row[8] or 0.0),
                        str(row[9]).strip() if row[9] else "",
                        str(row[10]).strip() if row[10] else "",
                        str(row[11]).strip() if row[11] else "",
                        str(row[12]).strip() if row[12] else "",
                        str(row[13]).strip() if row[13] else ""
                    ))
        conn.commit()
    except Exception as e:
        conn.rollback()
        print(f"Aviso durante a importação inicial: {e}")

def carregar_projetos():
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT nome FROM projetos ORDER BY id ASC")
        return [row[0] for row in cursor.fetchall()]


def _obter_projeto_id(cursor, nome_projeto):
    cursor.execute("SELECT id FROM projetos WHERE nome = ?", (nome_projeto,))
    res = cursor.fetchone()
    return res[0] if res else None


def listar_ciclos_abertos(nome_projeto="PROJETO AEV"):
    """Lista ciclos abertos do projeto, ordenados por ordem/id."""
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("""
            SELECT c.id, c.titulo, c.tipo, c.aberto, c.data_criacao, c.ordem
            FROM ciclos c
            JOIN projetos p ON c.projeto_id = p.id
            WHERE p.nome = ? AND c.aberto = 1
            ORDER BY c.ordem ASC, c.id ASC
        """, (nome_projeto,))
        return [dict(row) for row in cursor.fetchall()]


def contar_ciclos_abertos(nome_projeto="PROJETO AEV"):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT COUNT(*)
            FROM ciclos c
            JOIN projetos p ON c.projeto_id = p.id
            WHERE p.nome = ? AND c.aberto = 1
        """, (nome_projeto,))
        return cursor.fetchone()[0]


def obter_ciclo(ciclo_id):
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("""
            SELECT c.id, c.titulo, c.tipo, c.aberto, c.data_criacao, c.ordem, p.nome AS projeto
            FROM ciclos c
            JOIN projetos p ON c.projeto_id = p.id
            WHERE c.id = ?
        """, (ciclo_id,))
        row = cursor.fetchone()
        return dict(row) if row else None


def criar_ciclo(nome_projeto, tipo=TIPO_NOVA_MEDICAO, titulo=None):
    """
    Cria um novo ciclo aberto no projeto.
    Retorna o id do ciclo ou None em caso de erro/limite.
    """
    if tipo not in TIPOS_CICLO:
        tipo = TIPO_NOVA_MEDICAO
    if not titulo or not str(titulo).strip():
        titulo = TIPOS_CICLO.get(tipo, "Medição")

    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            projeto_id = _obter_projeto_id(cursor, nome_projeto)
            if not projeto_id:
                return None

            cursor.execute("""
                SELECT COUNT(*) FROM ciclos
                WHERE projeto_id = ? AND aberto = 1
            """, (projeto_id,))
            if cursor.fetchone()[0] >= LIMITE_ABAS_POR_PROJETO:
                return None

            cursor.execute("""
                SELECT COALESCE(MAX(ordem), -1) + 1 FROM ciclos
                WHERE projeto_id = ? AND aberto = 1
            """, (projeto_id,))
            ordem = cursor.fetchone()[0]

            from datetime import datetime
            agora = datetime.now().strftime("%d/%m/%Y %H:%M")
            cursor.execute("""
                INSERT INTO ciclos (projeto_id, titulo, tipo, aberto, data_criacao, ordem)
                VALUES (?, ?, ?, 1, ?, ?)
            """, (projeto_id, str(titulo).strip(), tipo, agora, ordem))
            conn.commit()
            return cursor.lastrowid
    except Exception as e:
        print(f"Erro ao criar ciclo: {e}")
        return None


def renomear_ciclo(ciclo_id, novo_titulo):
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "UPDATE ciclos SET titulo = ? WHERE id = ? AND aberto = 1",
                (str(novo_titulo).strip(), ciclo_id),
            )
            conn.commit()
            return cursor.rowcount > 0
    except Exception as e:
        print(f"Erro ao renomear ciclo: {e}")
        return False


def atualizar_ciclo(ciclo_id, titulo=None, tipo=None):
    """Atualiza título e/ou tipo de um ciclo aberto."""
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            if titulo is not None:
                cursor.execute(
                    "UPDATE ciclos SET titulo = ? WHERE id = ? AND aberto = 1",
                    (str(titulo).strip(), ciclo_id),
                )
            if tipo is not None and tipo in TIPOS_CICLO:
                cursor.execute(
                    "UPDATE ciclos SET tipo = ? WHERE id = ? AND aberto = 1",
                    (tipo, ciclo_id),
                )
            conn.commit()
            return True
    except Exception as e:
        print(f"Erro ao atualizar ciclo: {e}")
        return False


def excluir_ciclo(ciclo_id):
    """
    Remove a aba: apaga lançamentos do ciclo e marca como fechado.
    Não envia para histórico (diferente de fechar medição).
    """
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM lancamentos WHERE ciclo_id = ?", (ciclo_id,))
            cursor.execute("UPDATE ciclos SET aberto = 0 WHERE id = ?", (ciclo_id,))
            conn.commit()
            return True
    except Exception as e:
        print(f"Erro ao excluir ciclo: {e}")
        return False


def obter_ou_criar_ciclo_padrao(nome_projeto="PROJETO AEV"):
    """
    Retorna o id do primeiro ciclo aberto do projeto.
    Se não houver nenhum, cria um ciclo 'Medição' e retorna o id.
    Usado para manter compatibilidade enquanto a UI de abas não está ativa.
    """
    ciclos = listar_ciclos_abertos(nome_projeto)
    if ciclos:
        return ciclos[0]["id"]
    return criar_ciclo(nome_projeto, TIPO_NOVA_MEDICAO, "Medição")

def carregar_fornecedores():
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("""
        SELECT razao_social AS "RAZÃO SOCIAL", razao_fantasia AS "RAZÃO FANTASIA", 
               cnpj_cpf AS "CNPJ/CPF", banco AS "BANCO", agencia AS "AGÊNCIA", 
               conta AS "CONTA", tipo_pix AS "TIPO PIX", chave_pix AS "CHAVE PIX"
        FROM fornecedores
        ORDER BY razao_social ASC
        """)
        return [{k: (v if v is not None else "") for k, v in dict(row).items()} for row in cursor.fetchall()]

def adicionar_fornecedor(dados):
    """
    Valida e adiciona um novo fornecedor.
    Retorna (sucesso: bool, mensagem_erro: str). mensagem_erro vazia se ok.
    """
    if not dados or len(dados) < 3 or not dados[0] or not dados[2]:
        return False, (
            "Não foi possível cadastrar o prestador.\n\n"
            "Verifique se:\n"
            "• O campo Razão Social está preenchido\n"
            "• O campo CNPJ/CPF está preenchido"
        )

    razao_social = str(dados[0]).strip()
    cnpj_cpf = str(dados[2]).strip()

    try:
        with get_connection() as conn:
            cursor = conn.cursor()

            cursor.execute(
                "SELECT 1 FROM fornecedores WHERE razao_social = ? COLLATE NOCASE LIMIT 1",
                (razao_social,),
            )
            if cursor.fetchone():
                return False, (
                    "Não foi possível cadastrar o prestador.\n\n"
                    f"Já existe um cadastro com a Razão Social:\n« {razao_social} »\n\n"
                    "Verifique se:\n"
                    "• Esse prestador já não está na lista\n"
                    "• O nome não foi digitado com pequena diferença (acento, espaço)"
                )

            cursor.execute(
                "SELECT 1 FROM fornecedores WHERE cnpj_cpf = ? LIMIT 1",
                (cnpj_cpf,),
            )
            if cursor.fetchone():
                return False, (
                    "Não foi possível cadastrar o prestador.\n\n"
                    f"Já existe um cadastro com o CNPJ/CPF:\n« {cnpj_cpf} »\n\n"
                    "Verifique se:\n"
                    "• Esse prestador já não está cadastrado\n"
                    "• O CNPJ/CPF foi digitado corretamente"
                )

            cursor.execute("""
            INSERT INTO fornecedores
            (razao_social, razao_fantasia, cnpj_cpf, banco, agencia, conta, tipo_pix, chave_pix)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                razao_social,
                str(dados[1]).strip() if dados[1] else "",
                cnpj_cpf,
                str(dados[3]).strip() if dados[3] else "",
                str(dados[4]).strip() if dados[4] else "",
                str(dados[5]).strip() if dados[5] else "",
                str(dados[6]).strip() if dados[6] else "",
                str(dados[7]).strip() if dados[7] else ""
            ))
            conn.commit()
            return True, ""
    except sqlite3.IntegrityError:
        return False, (
            "Não foi possível cadastrar o prestador.\n\n"
            "Razão Social ou CNPJ/CPF já cadastrado.\n\n"
            "Verifique se esse prestador já existe na lista."
        )
    except Exception as e:
        print(f"Erro ao adicionar fornecedor: {e}")
        return False, (
            "Não foi possível cadastrar o prestador.\n\n"
            "Verifique se:\n"
            "• O banco de dados não está em uso por outro programa\n"
            "• Você tem permissão de escrita na pasta do sistema\n\n"
            f"Detalhe técnico: {e}"
        )


def atualizar_fornecedor(cnpj_original, dados):
    """
    Atualiza fornecedor existente.
    Retorna (sucesso: bool, mensagem_erro: str).
    """
    if not dados or len(dados) < 3 or not dados[0] or not dados[2]:
        return False, (
            "Não foi possível atualizar o prestador.\n\n"
            "Verifique se Razão Social e CNPJ/CPF estão preenchidos."
        )

    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            nova_razao = str(dados[0]).strip()
            novo_cnpj = str(dados[2]).strip()
            cnpj_orig = str(cnpj_original).strip()

            cursor.execute(
                """
                SELECT 1 FROM fornecedores
                WHERE razao_social = ? COLLATE NOCASE AND cnpj_cpf != ?
                LIMIT 1
                """,
                (nova_razao, cnpj_orig),
            )
            if cursor.fetchone():
                return False, (
                    "Não foi possível atualizar o prestador.\n\n"
                    f"Já existe outro cadastro com a Razão Social:\n« {nova_razao} »"
                )

            cursor.execute(
                """
                SELECT 1 FROM fornecedores
                WHERE cnpj_cpf = ? AND cnpj_cpf != ?
                LIMIT 1
                """,
                (novo_cnpj, cnpj_orig),
            )
            if cursor.fetchone():
                return False, (
                    "Não foi possível atualizar o prestador.\n\n"
                    f"Já existe outro cadastro com o CNPJ/CPF:\n« {novo_cnpj} »"
                )

            cursor.execute("""
            UPDATE fornecedores SET
                razao_social = ?, razao_fantasia = ?, cnpj_cpf = ?,
                banco = ?, agencia = ?, conta = ?, tipo_pix = ?, chave_pix = ?
            WHERE cnpj_cpf = ?
            """, (
                nova_razao,
                str(dados[1]).strip() if dados[1] else "",
                novo_cnpj,
                str(dados[3]).strip() if dados[3] else "",
                str(dados[4]).strip() if dados[4] else "",
                str(dados[5]).strip() if dados[5] else "",
                str(dados[6]).strip() if dados[6] else "",
                str(dados[7]).strip() if dados[7] else "",
                cnpj_orig,
            ))
            conn.commit()
            if cursor.rowcount > 0:
                return True, ""
            return False, (
                "Não foi possível atualizar o prestador.\n\n"
                "O registro original não foi encontrado.\n"
                "Tente selecionar o prestador novamente na lista."
            )
    except sqlite3.IntegrityError:
        return False, (
            "Não foi possível atualizar o prestador.\n\n"
            "Razão Social ou CNPJ/CPF conflita com outro cadastro já existente."
        )
    except Exception as e:
        print(f"Erro ao atualizar fornecedor: {e}")
        return False, (
            "Não foi possível atualizar o prestador.\n\n"
            f"Detalhe técnico: {e}"
        )

def carregar_lancamentos(nome_projeto="PROJETO AEV", apenas_ciclo_ativo=True, ciclo_id=None):
    """
    Carrega lançamentos do projeto.
    - Se ciclo_id for informado: só daquele ciclo.
    - Se apenas_ciclo_ativo e sem ciclo_id: usa o primeiro ciclo aberto (compatibilidade).
    """
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        query = f"""
        SELECT {SELECT_LANCAMENTO_COLUNAS}
        FROM lancamentos l
        JOIN projetos p ON l.projeto_id = p.id
        WHERE p.nome = ?
        """
        params = [nome_projeto]

        if ciclo_id is not None:
            query += " AND l.ciclo_id = ?"
            params.append(ciclo_id)
        elif apenas_ciclo_ativo:
            # Compatibilidade: restringe ao primeiro ciclo aberto do projeto
            cid = obter_ou_criar_ciclo_padrao(nome_projeto)
            if cid:
                query += " AND l.ciclo_id = ?"
                params.append(cid)

        query += " ORDER BY l.id ASC"
        cursor.execute(query, params)
        return [{k: (v if v is not None else "") for k, v in dict(row).items()} for row in cursor.fetchall()]


def buscar_lancamento_por_pc(pc, nome_projeto="PROJETO AEV", ciclo_id=None):
    """Busca lançamento pelo PC. Se ciclo_id informado, restringe ao ciclo."""
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        query = f"""
        SELECT {SELECT_LANCAMENTO_COLUNAS}
        FROM lancamentos l
        JOIN projetos p ON l.projeto_id = p.id
        WHERE p.nome = ? AND l.pc = ?
        """
        params = [nome_projeto, str(pc).strip()]
        if ciclo_id is not None:
            query += " AND l.ciclo_id = ?"
            params.append(ciclo_id)
        query += " ORDER BY l.id DESC LIMIT 1"
        cursor.execute(query, params)
        row = cursor.fetchone()
        if not row:
            return None
        return {k: (v if v is not None else "") for k, v in dict(row).items()}


def adicionar_lancamento(dados, nome_projeto="PROJETO AEV", ciclo_id=None):
    """
    Insere lançamento no ciclo informado (ou no ciclo padrão aberto do projeto).
    Retorna (sucesso: bool, mensagem_erro: str).
    """
    dados_limpos = _tratar_dados_lancamento(dados)
    if not dados_limpos[0]:
        return False, (
            "Não foi possível salvar o lançamento.\n\n"
            "Verifique se:\n"
            "• O campo PC está preenchido"
        )

    if ciclo_id is None:
        ciclo_id = obter_ou_criar_ciclo_padrao(nome_projeto)
    if not ciclo_id:
        return False, (
            "Não foi possível salvar o lançamento.\n\n"
            "Não há aba/ciclo ativo neste projeto.\n"
            "Crie uma nova aba ou selecione um projeto com medição aberta."
        )

    pc = dados_limpos[0]
    existente = buscar_lancamento_por_pc(pc, nome_projeto, ciclo_id=ciclo_id)
    if existente:
        return atualizar_lancamento(pc, dados, nome_projeto, ciclo_id=ciclo_id)

    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            projeto_id = _obter_projeto_id(cursor, nome_projeto)
            if not projeto_id:
                return False, (
                    "Não foi possível salvar o lançamento.\n\n"
                    f"O projeto « {nome_projeto} » não foi encontrado no sistema."
                )

            cursor.execute("""
            INSERT INTO lancamentos (
                projeto_id, ciclo_id, pc, prestador, atividade, valor_total, valor_liquido, valores_pagos,
                saldo_a_pagar, valor_a_pagar, saldo_aberto, status_pc, oc, end_id, nota_fiscal, regional
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (projeto_id, ciclo_id, *dados_limpos))
            conn.commit()
            return True, ""
    except sqlite3.IntegrityError:
        return False, (
            "Não foi possível salvar o lançamento.\n\n"
            f"Já pode existir um lançamento com o PC « {pc} » nesta aba.\n\n"
            "Verifique se:\n"
            "• O PC não está duplicado nesta medição\n"
            "• Você não está em outra aba do mesmo projeto"
        )
    except Exception as e:
        print(f"Erro ao adicionar lançamento: {e}")
        return False, (
            "Não foi possível salvar o lançamento.\n\n"
            "Verifique se:\n"
            "• O banco de dados não está em uso por outro programa\n"
            "• Há espaço em disco e permissão de escrita\n\n"
            f"Detalhe técnico: {e}"
        )


def atualizar_lancamento(pc_original, dados_atualizados, nome_projeto="PROJETO AEV", ciclo_id=None):
    """
    Atualiza lançamento. Se ciclo_id informado, restringe ao ciclo.
    Retorna (sucesso: bool, mensagem_erro: str).
    """
    dados_limpos = _tratar_dados_lancamento(dados_atualizados)
    novo_pc = dados_limpos[0]
    if not novo_pc:
        return False, (
            "Não foi possível atualizar o lançamento.\n\n"
            "Verifique se o campo PC está preenchido."
        )

    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            projeto_id = _obter_projeto_id(cursor, nome_projeto)
            if not projeto_id:
                return False, (
                    "Não foi possível atualizar o lançamento.\n\n"
                    f"O projeto « {nome_projeto} » não foi encontrado."
                )

            if ciclo_id is not None:
                cursor.execute("""
                UPDATE lancamentos SET
                    pc=?, prestador=?, atividade=?, valor_total=?, valor_liquido=?, valores_pagos=?,
                    saldo_a_pagar=?, valor_a_pagar=?, saldo_aberto=?, status_pc=?, oc=?, end_id=?,
                    nota_fiscal=?, regional=?
                WHERE pc=? AND projeto_id=? AND ciclo_id=?
                """, (*dados_limpos, str(pc_original).strip(), projeto_id, ciclo_id))
            else:
                cursor.execute("""
                UPDATE lancamentos SET
                    pc=?, prestador=?, atividade=?, valor_total=?, valor_liquido=?, valores_pagos=?,
                    saldo_a_pagar=?, valor_a_pagar=?, saldo_aberto=?, status_pc=?, oc=?, end_id=?,
                    nota_fiscal=?, regional=?
                WHERE pc=? AND projeto_id=?
                """, (*dados_limpos, str(pc_original).strip(), projeto_id))

                if cursor.rowcount == 0:
                    cursor.execute("""
                    UPDATE lancamentos SET
                        pc=?, prestador=?, atividade=?, valor_total=?, valor_liquido=?, valores_pagos=?,
                        saldo_a_pagar=?, valor_a_pagar=?, saldo_aberto=?, status_pc=?, oc=?, end_id=?,
                        nota_fiscal=?, regional=?
                    WHERE pc=? AND projeto_id=?
                    """, (*dados_limpos, str(novo_pc).strip(), projeto_id))

            conn.commit()
            if cursor.rowcount > 0:
                return True, ""
            return False, (
                "Não foi possível atualizar o lançamento.\n\n"
                f"O PC « {pc_original} » não foi encontrado nesta aba/projeto.\n\n"
                "Verifique se:\n"
                "• Você está na aba correta\n"
                "• O lançamento não foi excluído ou fechado"
            )
    except sqlite3.IntegrityError:
        return False, (
            "Não foi possível atualizar o lançamento.\n\n"
            "Há conflito de dados (possível PC duplicado nesta aba)."
        )
    except Exception as e:
        print(f"Erro ao atualizar lançamento: {e}")
        return False, (
            "Não foi possível atualizar o lançamento.\n\n"
            "Verifique se o banco de dados não está em uso por outro programa.\n\n"
            f"Detalhe técnico: {e}"
        )


def deletar_lancamento(pc, nome_projeto="PROJETO AEV", ciclo_id=None):
    """Retorna (sucesso: bool, mensagem_erro: str)."""
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            if ciclo_id is not None:
                cursor.execute("""
                DELETE FROM lancamentos
                WHERE pc = ? AND ciclo_id = ?
                  AND projeto_id = (SELECT id FROM projetos WHERE nome = ?)
                """, (str(pc).strip(), ciclo_id, nome_projeto))
            else:
                cursor.execute("""
                DELETE FROM lancamentos
                WHERE pc = ? AND projeto_id = (SELECT id FROM projetos WHERE nome = ?)
                """, (str(pc).strip(), nome_projeto))
            conn.commit()
            if cursor.rowcount > 0:
                return True, ""
            return False, (
                "Não foi possível excluir o lançamento.\n\n"
                f"O PC « {pc} » não foi encontrado nesta aba/projeto.\n\n"
                "Verifique se você está na aba correta ou se o registro já foi removido."
            )
    except Exception as e:
        print(f"Erro ao deletar lançamento: {e}")
        return False, (
            "Não foi possível excluir o lançamento.\n\n"
            "Verifique se o banco de dados não está em uso por outro programa.\n\n"
            f"Detalhe técnico: {e}"
        )

def exportar_planilha_medicao(caminho_saida="MEDICAO_FINANCEIRO.xlsx", nome_projeto="PROJETO AEV", ciclo_id=None):
    """
    Exporta a planilha financeira do ciclo.
    Retorna (sucesso: bool, mensagem_erro: str).
    """
    try:
        # Usa o caminho do template com get_resource_path
        template_path = get_resource_path(EXCEL_TEMPLATE)
        
        if not os.path.exists(template_path):
            print(f"Template {EXCEL_TEMPLATE} não encontrado em: {template_path}")
            return False, (
                "Não foi possível gerar a planilha.\n\n"
                "Verifique se:\n"
                f"• O arquivo modelo « {EXCEL_TEMPLATE} » está na pasta do sistema\n"
                "• A instalação do programa não está incompleta"
            )

        if ciclo_id is None:
            ciclo_id = obter_ou_criar_ciclo_padrao(nome_projeto)
        
        wb = openpyxl.load_workbook(template_path)
        
        with get_connection() as conn:
            cursor = conn.cursor()
            
            cursor.execute("""
            SELECT razao_social, razao_fantasia, cnpj_cpf, banco, agencia, conta, tipo_pix, chave_pix 
            FROM fornecedores
            """)
            fornecedores_rows = cursor.fetchall()
            
            dict_fornecedores = {}
            for f in fornecedores_rows:
                rs = str(f[0]).strip().upper() if f[0] else ""
                dict_fornecedores[rs] = {
                    "fantasia": str(f[1]).strip() if f[1] else "",
                    "cnpj_cpf": str(f[2]).strip() if f[2] else "",
                    "banco": str(f[3]).strip() if f[3] else "",
                    "tipo_pix": str(f[6]).strip() if f[6] else "",
                    "chave_pix": str(f[7]).strip() if f[7] else ""
                }

            ws_cad = _obter_aba(wb, ["CADASTROS DE FORNECEDORES", "CADASTRO"])
            if ws_cad:
                for r_idx, row in enumerate(fornecedores_rows, start=2):
                    for c_idx, val in enumerate(row, start=1):
                        ws_cad.cell(row=r_idx, column=c_idx, value=val)

            ws_lanc = _obter_aba(wb, ["27 07 2026", "LANÇAMENTOS"])
            if ws_lanc:
                ws_lanc.cell(row=1, column=1, value=f"{nome_projeto.upper()} - MEDIÇÃO")

                cursor.execute("""
                SELECT l.pc, l.prestador, l.atividade, l.valor_total, l.valor_liquido, l.valores_pagos,
                       l.saldo_a_pagar, l.valor_a_pagar, l.saldo_aberto, l.status_pc, l.oc, l.end_id,
                       l.nota_fiscal, l.regional
                FROM lancamentos l
                JOIN projetos p ON l.projeto_id = p.id
                WHERE p.nome = ? AND l.ciclo_id = ? AND l.valor_a_pagar > 0
                ORDER BY l.id ASC
                """, (nome_projeto, ciclo_id))
                lancamentos = cursor.fetchall()

                linha_inicio = 3
                total_registros = len(lancamentos)
                linha_subtotal_original = 8

                if total_registros > 5:
                    ws_lanc.insert_rows(linha_subtotal_original, amount=total_registros - 5)

                altura_modelo = ws_lanc.row_dimensions[linha_inicio].height or 45

                for idx in range(max(total_registros, 5)):
                    row_num = linha_inicio + idx
                    ws_lanc.row_dimensions[row_num].height = altura_modelo

                    if idx < total_registros:
                        row_data = lancamentos[idx]
                        template_row = linha_inicio + (idx % 5)

                        for c_idx, val in enumerate(row_data, start=1):
                            cell = ws_lanc.cell(row=row_num, column=c_idx)
                            cell.value = val

                            ref_cell = ws_lanc.cell(row=template_row, column=c_idx)
                            if ref_cell.has_style and row_num != template_row:
                                cell.font = copy(ref_cell.font)
                                cell.border = copy(ref_cell.border)
                                cell.fill = copy(ref_cell.fill)
                                cell.number_format = ref_cell.number_format
                                cell.alignment = copy(ref_cell.alignment)
                    else:
                        for c_idx in range(1, 15):
                            ws_lanc.cell(row=row_num, column=c_idx).value = None

                linha_subtotal_final = linha_inicio + max(total_registros, 5)
                ws_lanc.cell(row=linha_subtotal_final, column=9, value=f"=SUM(I{linha_inicio}:I{linha_subtotal_final-1})")
                ws_lanc.page_setup.print_area = None

            ws_tela = _obter_aba(wb, ["TELA AEV 27 07 2026", "TELA AEV"])
            if ws_tela:
                ws_tela.cell(row=1, column=1, value=f"TELA {nome_projeto.upper()} - TEMPLATE")

                cursor.execute("""
                SELECT 
                    l.prestador,
                    SUM(l.valor_a_pagar) AS total_a_pagar,
                    GROUP_CONCAT(DISTINCT NULLIF(TRIM(l.nota_fiscal), '')) AS nfs
                FROM lancamentos l
                JOIN projetos p ON l.projeto_id = p.id
                WHERE p.nome = ? AND l.ciclo_id = ? AND l.valor_a_pagar > 0
                  AND l.prestador IS NOT NULL AND l.prestador != ''
                GROUP BY l.prestador
                ORDER BY MIN(l.id) ASC
                """, (nome_projeto, ciclo_id))
                prestadores_pagamento = cursor.fetchall()

                linha_inicio_tela = 3
                total_prestadores = len(prestadores_pagamento)
                linha_total_original = 8

                if total_prestadores > 5:
                    ws_tela.insert_rows(linha_total_original, amount=total_prestadores - 5)

                altura_modelo_tela = ws_tela.row_dimensions[linha_inicio_tela].height or 45
                fmt_moeda = ws_tela.cell(row=3, column=4).number_format

                for idx in range(max(total_prestadores, 5)):
                    row_num = linha_inicio_tela + idx
                    ws_tela.row_dimensions[row_num].height = altura_modelo_tela

                    if idx < total_prestadores:
                        p_row = prestadores_pagamento[idx]
                        prestador_nome = str(p_row[0]).strip()
                        total_a_pagar = float(p_row[1] or 0.0)
                        nfs_raw = p_row[2]

                        nfs_str = ""
                        if nfs_raw:
                            nfs_lista = [n.strip() for n in str(nfs_raw).split(',') if n.strip()]
                            nfs_str = ", ".join(nfs_lista)

                        info_f = dict_fornecedores.get(prestador_nome.upper(), {})
                        template_row_tela = linha_inicio_tela + (idx % 5)

                        for c_idx in range(1, 9):
                            cell = ws_tela.cell(row=row_num, column=c_idx)
                            ref_cell = ws_tela.cell(row=template_row_tela, column=c_idx)
                            if ref_cell.has_style and row_num != template_row_tela:
                                cell.font = copy(ref_cell.font)
                                cell.border = copy(ref_cell.border)
                                cell.fill = copy(ref_cell.fill)
                                cell.alignment = copy(ref_cell.alignment)

                        ws_tela.cell(row=row_num, column=1, value=prestador_nome)
                        ws_tela.cell(row=row_num, column=2, value=info_f.get("fantasia", ""))
                        ws_tela.cell(row=row_num, column=3, value=info_f.get("cnpj_cpf", ""))
                        
                        cell_valor = ws_tela.cell(row=row_num, column=4, value=total_a_pagar)
                        cell_valor.number_format = fmt_moeda
                        
                        ws_tela.cell(row=row_num, column=5, value=nfs_str)
                        ws_tela.cell(row=row_num, column=6, value=info_f.get("tipo_pix", ""))
                        ws_tela.cell(row=row_num, column=7, value=info_f.get("banco", ""))
                        ws_tela.cell(row=row_num, column=8, value=info_f.get("chave_pix", ""))
                    else:
                        for c_idx in range(1, 9):
                            ws_tela.cell(row=row_num, column=c_idx).value = None

                linha_total_final = linha_inicio_tela + max(total_prestadores, 5)
                cell_tot = ws_tela.cell(row=linha_total_final, column=4, value=f"=SUM(D{linha_inicio_tela}:D{linha_total_final-1})")
                cell_tot.number_format = fmt_moeda

                ws_tela.page_setup.print_area = None
                if "Print_Area" in ws_tela.defined_names:
                    del ws_tela.defined_names["Print_Area"]

        try:
            wb.save(caminho_saida)
        except PermissionError:
            return False, (
                "Não foi possível salvar a planilha.\n\n"
                "Verifique se:\n"
                "• O arquivo não está aberto no Excel (ou outro programa)\n"
                "• Você fechou a planilha antes de exportar de novo\n"
                "• A pasta de destino permite gravação\n"
                "• O arquivo não está marcado como somente leitura"
            )
        except OSError as e:
            # Windows: arquivo em uso costuma vir como OSError / WinError 32
            erro_txt = str(e).lower()
            if "permission" in erro_txt or "being used" in erro_txt or "em uso" in erro_txt or getattr(e, "winerror", None) == 32:
                return False, (
                    "Não foi possível salvar a planilha.\n\n"
                    "O arquivo parece estar aberto ou bloqueado.\n\n"
                    "Verifique se:\n"
                    "• A planilha não está aberta no Excel\n"
                    "• Você fechou o arquivo e tentou exportar novamente\n"
                    "• Escolheu outro nome/local se precisar manter a versão antiga aberta"
                )
            return False, (
                "Não foi possível salvar a planilha.\n\n"
                "Verifique se:\n"
                "• O caminho de destino é válido\n"
                "• Há espaço em disco e permissão de escrita\n\n"
                f"Detalhe técnico: {e}"
            )

        return True, ""
    except PermissionError:
        return False, (
            "Não foi possível salvar a planilha.\n\n"
            "Verifique se o arquivo não está aberto no Excel e tente novamente."
        )
    except Exception as e:
        print(f"Erro ao exportar planilha para o Financeiro: {e}")
        erro_txt = str(e).lower()
        if "permission" in erro_txt or "being used" in erro_txt or getattr(e, "winerror", None) == 32:
            return False, (
                "Não foi possível salvar a planilha.\n\n"
                "Verifique se:\n"
                "• O arquivo não está aberto no Excel\n"
                "• Você fechou a planilha antes de sobrescrever"
            )
        return False, (
            "Não foi possível gerar a planilha.\n\n"
            "Verifique se:\n"
            "• O template do sistema está disponível\n"
            "• Há lançamentos no ciclo atual\n"
            "• O destino escolhido permite gravação\n\n"
            f"Detalhe técnico: {e}"
        )

def fechar_e_consolidar_medicao(caminho_exportacao, nome_projeto="PROJETO AEV", ciclo_id=None):
    """
    Consolida um ciclo de medição:
    1. Gera a planilha Excel do período (só do ciclo).
    2. Transfere os dados do ciclo para o histórico (com tipo_ciclo).
    3. Atualiza os saldos dos lançamentos daquele ciclo.
    4. Marca o ciclo como fechado (aberto = 0).
    """
    if ciclo_id is None:
        ciclo_id = obter_ou_criar_ciclo_padrao(nome_projeto)
    if not ciclo_id:
        return False, (
            "Não foi possível fechar a medição.\n\n"
            "Não há ciclo/aba aberta neste projeto."
        )

    sucesso_exportacao, msg_export = exportar_planilha_medicao(
        caminho_exportacao, nome_projeto, ciclo_id=ciclo_id
    )
    if not sucesso_exportacao:
        return False, (
            "Não foi possível fechar a medição porque a planilha não pôde ser salva.\n"
            "Nenhum dado do banco foi alterado.\n\n"
            + (msg_export or "Tente fechar o arquivo no Excel e repetir.")
        )

    try:
        with get_connection() as conn:
            cursor = conn.cursor()

            from datetime import datetime
            mes_ref = datetime.now().strftime("%m/%Y")
            data_fechamento = datetime.now().strftime("%d/%m/%Y")

            # Tipo do ciclo para o histórico
            cursor.execute("SELECT tipo FROM ciclos WHERE id = ?", (ciclo_id,))
            row_tipo = cursor.fetchone()
            tipo_ciclo = row_tipo[0] if row_tipo else TIPO_NOVA_MEDICAO

            cursor.execute("""
            INSERT INTO historico_medicoes
            (projeto_id, data_fechamento, mes_referencia, pc, prestador, atividade,
             valor_total, valor_liquido, valores_pagos, saldo_a_pagar, valor_a_pagar, saldo_aberto,
             status_pc, oc, end_id, nota_fiscal, regional, valor_medido, status, tipo_ciclo)
            SELECT
                p.id, ?, ?, l.pc, l.prestador, l.atividade,
                l.valor_total, l.valor_liquido, l.valores_pagos, l.saldo_a_pagar, l.valor_a_pagar, l.saldo_aberto,
                l.status_pc, l.oc, l.end_id, l.nota_fiscal, l.regional, l.valor_a_pagar, l.status_pc, ?
            FROM lancamentos l
            JOIN projetos p ON l.projeto_id = p.id
            WHERE p.nome = ? AND l.ciclo_id = ? AND l.valor_a_pagar > 0
            """, (data_fechamento, mes_ref, tipo_ciclo, nome_projeto, ciclo_id))

            cursor.execute("""
            SELECT l.id, l.valor_total, l.valores_pagos, l.valor_a_pagar
            FROM lancamentos l
            JOIN projetos p ON l.projeto_id = p.id
            WHERE p.nome = ? AND l.ciclo_id = ?
            """, (nome_projeto, ciclo_id))

            registros = cursor.fetchall()

            for reg in registros:
                l_id, v_total, v_pagos, v_a_pagar = reg

                novos_pagos = (v_pagos or 0.0) + (v_a_pagar or 0.0)
                novo_saldo_a_pagar = max(0.0, (v_total or 0.0) - novos_pagos)
                novo_valor_a_pagar = 0.0
                novo_saldo_aberto = novo_saldo_a_pagar

                cursor.execute("""
                UPDATE lancamentos SET
                    valores_pagos = ?,
                    saldo_a_pagar = ?,
                    valor_a_pagar = ?,
                    saldo_aberto = ?
                WHERE id = ?
                """, (novos_pagos, novo_saldo_a_pagar, novo_valor_a_pagar, novo_saldo_aberto, l_id))

            # Marca o ciclo como fechado
            cursor.execute("UPDATE ciclos SET aberto = 0 WHERE id = ?", (ciclo_id,))

            conn.commit()
            return True, ""
    except Exception as e:
        print(f"Erro ao consolidar medição no banco de dados: {e}")
        return False, (
            "A planilha foi gerada, mas houve erro ao consolidar no banco.\n\n"
            "Verifique se o banco de dados não está em uso por outro programa.\n\n"
            f"Detalhe técnico: {e}"
        )


def limpar_lancamentos_projeto(nome_projeto="PROJETO AEV", ciclo_id=None):
    """
    Remove lançamentos do projeto.
    Se ciclo_id informado: só daquele ciclo (e fecha o ciclo).
    Caso contrário: todos os lançamentos e ciclos abertos do projeto.
    Retorna (sucesso: bool, mensagem_erro: str).
    """
    try:
        with get_connection() as conn:
            cursor = conn.cursor()
            projeto_id = _obter_projeto_id(cursor, nome_projeto)
            if not projeto_id:
                return False, (
                    "Não foi possível limpar os lançamentos.\n\n"
                    f"O projeto « {nome_projeto} » não foi encontrado."
                )

            if ciclo_id is not None:
                cursor.execute(
                    "DELETE FROM lancamentos WHERE projeto_id = ? AND ciclo_id = ?",
                    (projeto_id, ciclo_id),
                )
                cursor.execute(
                    "UPDATE ciclos SET aberto = 0 WHERE id = ?",
                    (ciclo_id,),
                )
            else:
                cursor.execute(
                    "DELETE FROM lancamentos WHERE projeto_id = ?",
                    (projeto_id,),
                )
                cursor.execute(
                    "UPDATE ciclos SET aberto = 0 WHERE projeto_id = ? AND aberto = 1",
                    (projeto_id,),
                )

            conn.commit()
            return True, ""
    except Exception as e:
        print(f"Erro ao limpar lançamentos do projeto '{nome_projeto}': {e}")
        return False, (
            "Não foi possível limpar os lançamentos.\n\n"
            "Verifique se:\n"
            "• O banco de dados não está em uso por outro programa\n"
            "• Há permissão de escrita na pasta do sistema\n\n"
            f"Detalhe técnico: {e}"
        )

def rotulo_tipo_ciclo(tipo):
    """Nome amigável do tipo de ciclo para exibição."""
    if not tipo:
        return "Medição"
    return TIPOS_CICLO.get(tipo, "Medição")


def carregar_historico_medicoes(nome_projeto="PROJETO AEV", mes_referencia=None, tipo_ciclo=None):
    """Carrega a lista de medições fechadas do projeto, agrupadas por data + tipo."""
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        query = """
        SELECT
            data_fechamento,
            COALESCE(tipo_ciclo, 'NOVA_MEDICAO') AS tipo_ciclo,
            SUM(valor_medido) as total_pago,
            COUNT(*) as quantidade_pcs
        FROM historico_medicoes h
        JOIN projetos p ON h.projeto_id = p.id
        WHERE p.nome = ?
        """
        params = [nome_projeto]

        if mes_referencia:
            query += " AND h.mes_referencia = ?"
            params.append(mes_referencia)

        if tipo_ciclo:
            query += " AND COALESCE(h.tipo_ciclo, 'NOVA_MEDICAO') = ?"
            params.append(tipo_ciclo)

        query += """
            GROUP BY data_fechamento, COALESCE(tipo_ciclo, 'NOVA_MEDICAO')
            ORDER BY data_fechamento DESC, tipo_ciclo ASC
        """

        cursor.execute(query, params)
        return [dict(row) for row in cursor.fetchall()]


def carregar_detalhes_medicao(data_fechamento, nome_projeto="PROJETO AEV", tipo_ciclo=None):
    """Carrega os detalhes de uma medição específica (data + tipo opcional)."""
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        query = """
        SELECT
            h.pc AS "PC",
            h.prestador AS "PRESTADOR",
            COALESCE(h.atividade, '') AS "ATIVIDADE/SERVIÇO",
            COALESCE(h.valor_total, 0) AS "VALOR TOTAL DO PC",
            COALESCE(h.valor_liquido, 0) AS "VALOR LIQUIDO",
            COALESCE(h.valores_pagos, 0) AS "VALORES PAGOS",
            COALESCE(h.saldo_a_pagar, 0) AS "SALDO A PAGAR",
            COALESCE(h.valor_a_pagar, h.valor_medido, 0) AS "VALOR A PAGAR",
            COALESCE(h.saldo_aberto, 0) AS "SALDO EM ABERTO",
            COALESCE(h.status_pc, h.status, '') AS "STATUS PC",
            COALESCE(h.oc, '') AS "OC",
            COALESCE(h.end_id, '') AS "END ID",
            COALESCE(h.nota_fiscal, '') AS "NOTA FISCAL",
            COALESCE(h.regional, '') AS "REGIONAL"
        FROM historico_medicoes h
        JOIN projetos p ON h.projeto_id = p.id
        WHERE p.nome = ? AND h.data_fechamento = ?
        """
        params = [nome_projeto, data_fechamento]
        if tipo_ciclo:
            query += " AND COALESCE(h.tipo_ciclo, 'NOVA_MEDICAO') = ?"
            params.append(tipo_ciclo)
        query += " ORDER BY h.id ASC"

        cursor.execute(query, params)
        return [dict(row) for row in cursor.fetchall()]


def carregar_estatisticas_por_tipo(nome_projeto="PROJETO AEV", tipo_ciclo=None, prestador_filtro=None):
    """
    Estatísticas a partir do histórico (pagamentos já fechados), filtráveis por tipo.
    Útil para: quanto de emergencial, saldo final, etc.
    """
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        query_tot = """
        SELECT
            COALESCE(SUM(h.valor_medido), 0) AS total_pago,
            COUNT(*) AS quantidade_pcs,
            COUNT(DISTINCT h.data_fechamento) AS quantidade_fechamentos
        FROM historico_medicoes h
        JOIN projetos p ON h.projeto_id = p.id
        WHERE p.nome = ?
        """
        params = [nome_projeto]
        if tipo_ciclo:
            query_tot += " AND COALESCE(h.tipo_ciclo, 'NOVA_MEDICAO') = ?"
            params.append(tipo_ciclo)
        if prestador_filtro:
            query_tot += " AND h.prestador = ?"
            params.append(prestador_filtro)

        cursor.execute(query_tot, params)
        totais = dict(cursor.fetchone())

        query_prest = """
        SELECT
            h.prestador,
            COALESCE(SUM(h.valor_medido), 0) AS pago,
            COUNT(*) AS qtd_pcs
        FROM historico_medicoes h
        JOIN projetos p ON h.projeto_id = p.id
        WHERE p.nome = ? AND h.prestador IS NOT NULL AND h.prestador != ''
        """
        params_p = [nome_projeto]
        if tipo_ciclo:
            query_prest += " AND COALESCE(h.tipo_ciclo, 'NOVA_MEDICAO') = ?"
            params_p.append(tipo_ciclo)
        if prestador_filtro:
            query_prest += " AND h.prestador = ?"
            params_p.append(prestador_filtro)
        query_prest += " GROUP BY h.prestador ORDER BY pago DESC"

        cursor.execute(query_prest, params_p)
        prestadores = [dict(row) for row in cursor.fetchall()]

        return {"totais": totais, "prestadores": prestadores}

def carregar_meses_com_historico(nome_projeto="PROJETO AEV"):
    """Carrega os meses que têm medições fechadas"""
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
        SELECT DISTINCT mes_referencia
        FROM historico_medicoes h
        JOIN projetos p ON h.projeto_id = p.id
        WHERE p.nome = ?
        ORDER BY mes_referencia DESC
        """, (nome_projeto,))
        return [row[0] for row in cursor.fetchall()]

def carregar_estatisticas_projeto(nome_projeto="PROJETO AEV", prestador_filtro=None):
    """Carrega estatísticas do projeto (inclui PCs quitados para totais e pagamentos)."""
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        query_total = """
        SELECT 
            SUM(l.valor_total) as total_projeto,
            SUM(l.valores_pagos) as total_pago,
            SUM(l.saldo_aberto) as saldo_em_aberto
        FROM lancamentos l
        JOIN projetos p ON l.projeto_id = p.id
        WHERE p.nome = ?
        """
        params = [nome_projeto]
        
        if prestador_filtro:
            query_total += " AND l.prestador = ?"
            params.append(prestador_filtro)
            
        cursor.execute(query_total, params)
        totais = dict(cursor.fetchone())
        
        query_prestadores = """
        SELECT 
            l.prestador,
            SUM(l.valores_pagos) as pago,
            SUM(l.saldo_aberto) as saldo,
            SUM(l.valor_total) as total
        FROM lancamentos l
        JOIN projetos p ON l.projeto_id = p.id
        WHERE p.nome = ? AND l.prestador IS NOT NULL AND l.prestador != ''
        """
        params_prest = [nome_projeto]
        
        if prestador_filtro:
            query_prestadores += " AND l.prestador = ?"
            params_prest.append(prestador_filtro)
            
        query_prestadores += " GROUP BY l.prestador ORDER BY total DESC"
        
        cursor.execute(query_prestadores, params_prest)
        prestadores = [dict(row) for row in cursor.fetchall()]
        
        return {
            "totais": totais,
            "prestadores": prestadores
        }

def carregar_lancamentos_prestador(nome_projeto="PROJETO AEV", prestador=None):
    """Carrega os lançamentos de um prestador específico"""
    if not prestador:
        return []
        
    with get_connection() as conn:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute(f"""
        SELECT {SELECT_LANCAMENTO_COLUNAS}
        FROM lancamentos l
        JOIN projetos p ON l.projeto_id = p.id
        WHERE p.nome = ? AND l.prestador = ?
        ORDER BY l.saldo_aberto DESC, l.id ASC
        """, (nome_projeto, prestador))
        
        return [dict(row) for row in cursor.fetchall()]

if __name__ == "__main__":
    inicializar_banco()
    print("Banco inicializado e validado com sucesso!")